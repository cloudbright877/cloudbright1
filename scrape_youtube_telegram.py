"""
Скрипт для сбора Telegram-контактов с YouTube каналов.
Читает YouTubers_Report_v2_CLEANED.xlsx, парсит:
  1) About-страницу канала (ссылки из ytInitialData)
  2) Описания последних 5-7 видео

Ищет t.me / telegram.me ссылки.

Запуск: python scrape_youtube_telegram.py [--skip-existing] [--start N] [--limit N]
Результат: YouTubers_Report_v3_ENRICHED.xlsx + scrape_youtube.log
Прогресс сохраняется в scrape_progress.json (можно прервать и продолжить).
"""

import re
import json
import time
import random
import subprocess
import sys
import os
import logging
import argparse
import tempfile
from pathlib import Path
from datetime import datetime

import openpyxl

# --- Config ---
INPUT_FILE = "YouTubers_Report_v2_CLEANED.xlsx"
OUTPUT_FILE = "YouTubers_Report_v3_ENRICHED.xlsx"
PROGRESS_FILE = "scrape_progress.json"
MAX_VIDEOS = 7
DELAY_CHANNEL = (3.0, 6.0)  # Delay between channels
DELAY_VIDEO = (1.5, 3.5)    # Delay between video fetches

# Columns (0-indexed)
COL_NUM = 0
COL_CHANNEL = 2
COL_SUBS = 3
COL_URL = 4
COL_TELEGRAM = 16
COL_NOTES = 23

# --- Logging ---
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[
        logging.FileHandler("scrape_youtube.log", encoding="utf-8"),
        logging.StreamHandler(sys.stdout),
    ],
)
log = logging.getLogger(__name__)

# TG link blacklist (bots, generic pages)
TG_BLACKLIST = {
    "share", "botfather", "stickers", "gif", "youtube", "addstickers",
    "proxy", "socks", "iv", "joinchat",
}


def normalize_url(url: str) -> str:
    url = url.strip()
    if not url.startswith("http"):
        url = "https://" + url
    return url


def extract_channel_handle(url: str) -> str | None:
    """Extract @handle from channel URL if available."""
    m = re.search(r'youtube\.com/@([A-Za-z0-9_\-]+)', url)
    return m.group(1) if m else None


def extract_channel_id(url: str) -> str | None:
    """Extract channel ID from URL."""
    m = re.search(r'youtube\.com/channel/(UC[A-Za-z0-9_\-]+)', url)
    return m.group(1) if m else None


def run_ytdlp(args: list[str], timeout: int = 60) -> tuple[str, bool, str]:
    """Run yt-dlp and return (stdout, success, stderr)."""
    cmd = [sys.executable, "-m", "yt_dlp"] + args
    try:
        result = subprocess.run(
            cmd, capture_output=True, text=True,
            timeout=timeout, encoding="utf-8", errors="replace",
        )
        return result.stdout, result.returncode == 0, result.stderr
    except subprocess.TimeoutExpired:
        return "", False, "TIMEOUT"
    except Exception as e:
        return "", False, str(e)


def extract_tg_links(text: str) -> list[str]:
    """Extract and clean Telegram links from text."""
    if not text:
        return []

    found = set()

    # t.me/... links
    for m in re.finditer(r'(?:https?://)?t\.me/(\+?[A-Za-z0-9_\-/]+)', text, re.I):
        raw = m.group(0).rstrip("/.,;:!?)")
        username = m.group(1).rstrip("/.,;:!?)").lower().split("/")[0]
        if username in TG_BLACKLIST or len(username) < 3:
            continue
        if not raw.startswith("http"):
            raw = "https://" + raw
        found.add(raw)

    # telegram.me/... links
    for m in re.finditer(r'(?:https?://)?telegram\.me/(\+?[A-Za-z0-9_\-/]+)', text, re.I):
        raw = m.group(0).rstrip("/.,;:!?)")
        username = m.group(1).rstrip("/.,;:!?)").lower().split("/")[0]
        if username in TG_BLACKLIST or len(username) < 3:
            continue
        link = re.sub(r'telegram\.me', 't.me', raw)
        if not link.startswith("http"):
            link = "https://" + link
        found.add(link)

    return sorted(found)


def resolve_channel_handle(channel_url: str) -> str | None:
    """Resolve @handle for a channel URL by checking first video metadata."""
    url = normalize_url(channel_url)
    out, ok, _ = run_ytdlp([
        "--skip-download", "--no-warnings",
        "--print", "%(uploader_id)s",
        "--playlist-items", "1",
        url,
    ], timeout=45)
    if ok and out.strip():
        handle = out.strip().lstrip("@")
        if handle and handle != "NA":
            return handle
    return None


def get_about_page_links(channel_url: str) -> list[str]:
    """
    Download channel /about page via yt-dlp --write-pages,
    parse ytInitialData JSON for t.me links.
    yt-dlp saves .dump files to cwd, so we cd to a temp dir.
    """
    url = normalize_url(channel_url)

    # Determine about URL - prefer @handle (works better for about page)
    handle = extract_channel_handle(url)
    if not handle:
        handle = resolve_channel_handle(url)

    if handle:
        about_url = f"https://www.youtube.com/@{handle}/about"
    else:
        ch_id = extract_channel_id(url)
        if ch_id:
            about_url = f"https://www.youtube.com/channel/{ch_id}/about"
        else:
            about_url = url.rstrip("/") + "/about"

    # yt-dlp --write-pages saves .dump to cwd, so use temp dir as cwd
    with tempfile.TemporaryDirectory() as tmpdir:
        args = [
            "--skip-download", "--no-warnings",
            "--write-pages",
            "--playlist-items", "0",
            about_url,
        ]
        # Run yt-dlp with cwd set to tmpdir
        cmd = [sys.executable, "-m", "yt_dlp"] + args
        try:
            result = subprocess.run(
                cmd, capture_output=True, text=True,
                timeout=45, encoding="utf-8", errors="replace",
                cwd=tmpdir,
            )
            ok = result.returncode == 0
        except subprocess.TimeoutExpired:
            log.warning("    About page timeout")
            return []
        except Exception as e:
            log.warning("    About page error: %s", e)
            return []

        # Find dump file in tmpdir (where yt-dlp saved it)
        dump_files = list(Path(tmpdir).glob("*.dump"))
        if not dump_files:
            log.warning("    No dump file created")
            return []

        content = dump_files[0].read_text(encoding="utf-8", errors="replace")

    # Extract ytInitialData
    match = re.search(r'var ytInitialData\s*=\s*({.*?});</script>', content, re.DOTALL)
    if not match:
        return extract_tg_links(content)

    try:
        data_text = json.dumps(json.loads(match.group(1)))
    except json.JSONDecodeError:
        data_text = match.group(1)

    return extract_tg_links(data_text)


def get_video_descriptions(channel_url: str, max_videos: int = 7) -> list[dict]:
    """Get video IDs, then fetch each description individually."""
    url = normalize_url(channel_url)
    videos_url = url.rstrip("/") + "/videos"

    # Get video list
    args = [
        "--skip-download", "--no-warnings", "--flat-playlist",
        "--print", '{"id":"%(id)s","title":"%(title)s"}',
        "--playlist-items", f"1:{max_videos}",
        videos_url,
    ]
    stdout, ok, _ = run_ytdlp(args, timeout=60)

    if not ok or not stdout.strip():
        # Fallback without /videos
        args[-1] = url
        stdout, ok, _ = run_ytdlp(args, timeout=60)

    if not ok or not stdout.strip():
        return []

    videos = []
    for line in stdout.strip().split("\n"):
        line = line.strip()
        if not line:
            continue
        try:
            videos.append(json.loads(line))
        except json.JSONDecodeError:
            continue

    # Fetch individual descriptions
    results = []
    for vid in videos:
        vid_id = vid.get("id", "")
        if not vid_id or vid_id == "NA":
            continue

        time.sleep(random.uniform(*DELAY_VIDEO))

        vid_url = f"https://www.youtube.com/watch?v={vid_id}"
        out, ok2, _ = run_ytdlp([
            "--skip-download", "--no-warnings", "--dump-json", vid_url
        ], timeout=40)

        desc = ""
        if ok2 and out.strip():
            try:
                data = json.loads(out)
                desc = data.get("description", "")
            except json.JSONDecodeError:
                pass

        results.append({
            "id": vid_id,
            "title": vid.get("title", ""),
            "description": desc,
        })

    return results


def scrape_channel(channel_url: str, channel_name: str) -> dict:
    """Full scrape of a single channel. Returns found TG links + metadata."""
    result = {
        "telegram_links": [],
        "sources": [],
        "errors": [],
    }

    # 1) About page links
    log.info("  [1/2] Scraping about page...")
    try:
        about_links = get_about_page_links(channel_url)
        if about_links:
            log.info("    Found %d TG link(s) in about page: %s", len(about_links), about_links)
            result["telegram_links"].extend(about_links)
            result["sources"].extend([f"about: {l}" for l in about_links])
        else:
            log.info("    No TG links on about page")
    except Exception as e:
        log.warning("    About page error: %s", e)
        result["errors"].append(f"about: {e}")

    time.sleep(random.uniform(1.5, 3.0))

    # 2) Video descriptions
    log.info("  [2/2] Scraping last %d video descriptions...", MAX_VIDEOS)
    try:
        videos = get_video_descriptions(channel_url, MAX_VIDEOS)
        log.info("    Got %d video(s)", len(videos))

        for i, vid in enumerate(videos):
            vid_links = extract_tg_links(vid["description"])
            if vid_links:
                log.info("    Video %d '%s': %s", i + 1, vid["title"][:40], vid_links)
                result["telegram_links"].extend(vid_links)
                result["sources"].extend([f'video "{vid["title"][:30]}": {l}' for l in vid_links])
    except Exception as e:
        log.warning("    Videos error: %s", e)
        result["errors"].append(f"videos: {e}")

    # Deduplicate
    result["telegram_links"] = sorted(set(result["telegram_links"]))
    return result


def load_progress() -> dict:
    path = Path(PROGRESS_FILE)
    if path.exists():
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    return {"completed": {}}


def save_progress(progress: dict):
    with open(PROGRESS_FILE, "w", encoding="utf-8") as f:
        json.dump(progress, f, ensure_ascii=False, indent=2)


def main():
    parser = argparse.ArgumentParser(description="Scrape YouTube channels for Telegram contacts")
    parser.add_argument("--skip-existing", action="store_true",
                        help="Skip channels that already have Telegram in spreadsheet")
    parser.add_argument("--start", type=int, default=1,
                        help="Start from channel number N (1-based)")
    parser.add_argument("--limit", type=int, default=0,
                        help="Process only N channels (0 = all)")
    parser.add_argument("--reset", action="store_true",
                        help="Reset progress file and start fresh")
    args = parser.parse_args()

    log.info("=" * 60)
    log.info("YouTube -> Telegram Scraper")
    log.info("Input: %s | Output: %s", INPUT_FILE, OUTPUT_FILE)
    log.info("Options: skip_existing=%s start=%d limit=%d",
             args.skip_existing, args.start, args.limit)
    log.info("=" * 60)

    # Load workbook
    wb = openpyxl.load_workbook(INPUT_FILE)
    ws = wb.active

    # Progress
    if args.reset and Path(PROGRESS_FILE).exists():
        os.remove(PROGRESS_FILE)
        log.info("Progress reset")
    progress = load_progress()
    completed = progress.get("completed", {})

    # Collect channels
    channels = []
    for row_idx in range(2, ws.max_row + 1):
        num = ws.cell(row=row_idx, column=COL_NUM + 1).value
        name = ws.cell(row=row_idx, column=COL_CHANNEL + 1).value
        url = ws.cell(row=row_idx, column=COL_URL + 1).value
        subs = ws.cell(row=row_idx, column=COL_SUBS + 1).value
        existing_tg = ws.cell(row=row_idx, column=COL_TELEGRAM + 1).value

        if name and str(name).startswith("---"):
            continue
        if num is None or not url:
            continue

        channels.append({
            "row": row_idx,
            "num": int(num) if isinstance(num, (int, float)) else num,
            "name": str(name or ""),
            "url": str(url),
            "subs": str(subs or ""),
            "existing_tg": str(existing_tg) if existing_tg else "",
        })

    log.info("Total channels in spreadsheet: %d", len(channels))

    # Apply filters
    if args.start > 1:
        channels = [c for c in channels if c["num"] >= args.start]
    if args.skip_existing:
        before = len(channels)
        channels = [c for c in channels if not c["existing_tg"]]
        log.info("After skip-existing filter: %d (removed %d)", len(channels), before - len(channels))
    if args.limit > 0:
        channels = channels[:args.limit]

    log.info("Channels to process: %d (already completed: %d)",
             len(channels), len(completed))

    stats = {"processed": 0, "found_new": 0, "updated": 0, "errors": 0, "skipped": 0}

    for i, ch in enumerate(channels):
        row_key = str(ch["num"])

        if row_key in completed:
            stats["skipped"] += 1
            continue

        log.info("")
        log.info("=" * 50)
        log.info("[%d/%d] #%s %s (%s subs)",
                 i + 1, len(channels), ch["num"], ch["name"], ch["subs"])
        log.info("  URL: %s", ch["url"])
        if ch["existing_tg"]:
            log.info("  Existing TG: %s", ch["existing_tg"])

        result = scrape_channel(ch["url"], ch["name"])
        stats["processed"] += 1

        if result["errors"]:
            stats["errors"] += 1

        if result["telegram_links"]:
            new_tg_str = "; ".join(result["telegram_links"])
            log.info("  >>> FOUND: %s", new_tg_str)

            existing = ch["existing_tg"]
            if existing:
                existing_lower = existing.lower()
                truly_new = []
                for link in result["telegram_links"]:
                    link_name = link.split("/")[-1].lower().split("?")[0]
                    if link.lower() not in existing_lower and link_name not in existing_lower:
                        truly_new.append(link)

                if truly_new:
                    merged = existing + "; " + "; ".join(truly_new)
                    ws.cell(row=ch["row"], column=COL_TELEGRAM + 1).value = merged
                    stats["found_new"] += 1
                    stats["updated"] += 1
                    log.info("  MERGED: %s", merged)
                else:
                    log.info("  All links already known")
            else:
                ws.cell(row=ch["row"], column=COL_TELEGRAM + 1).value = new_tg_str
                stats["found_new"] += 1
                stats["updated"] += 1
                log.info("  NEW: %s", new_tg_str)

            # Source annotation in Notes
            src_short = "; ".join(result["sources"][:3])
            note_tag = f" [TG found {datetime.now().strftime('%m/%d')}: {src_short}]"
            current_notes = ws.cell(row=ch["row"], column=COL_NOTES + 1).value or ""
            if "[TG found" not in str(current_notes):
                ws.cell(row=ch["row"], column=COL_NOTES + 1).value = str(current_notes) + note_tag
        else:
            log.info("  No Telegram links found")

        # Save progress
        completed[row_key] = {
            "name": ch["name"],
            "found": result["telegram_links"],
            "ts": datetime.now().isoformat(),
        }
        save_progress({"completed": completed})

        # Checkpoint save every 10 channels
        if stats["processed"] % 10 == 0:
            wb.save(OUTPUT_FILE)
            log.info("  [Checkpoint] Saved %s (%d processed, %d found)",
                     OUTPUT_FILE, stats["processed"], stats["found_new"])

        # Delay between channels
        delay = random.uniform(*DELAY_CHANNEL)
        log.info("  Waiting %.1fs...", delay)
        time.sleep(delay)

    # Final save
    wb.save(OUTPUT_FILE)

    log.info("")
    log.info("=" * 60)
    log.info("FINISHED!")
    log.info("  Processed: %d channels", stats["processed"])
    log.info("  Skipped (already done): %d", stats["skipped"])
    log.info("  Found new TG contacts: %d", stats["found_new"])
    log.info("  Updated rows: %d", stats["updated"])
    log.info("  Errors: %d", stats["errors"])
    log.info("  Output: %s", OUTPUT_FILE)
    log.info("=" * 60)


if __name__ == "__main__":
    main()
