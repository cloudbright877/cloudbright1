"""
Cleanup and enrich YouTubers_Report_v2_FULL.xlsx
- Remove duplicate/scored rows ([8.0], [AD:xx] prefixes)
- Remove broken rows (LINE 48, empty names)
- Remove URL duplicates (keep the one with most data)
- Re-number rows
- Add new channels
- Enrich existing Uniminepool entries with research data
"""

import openpyxl
import re
import copy
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def load_row_data(ws, row, headers):
    """Load a row as a dictionary"""
    data = {}
    for col in range(1, len(headers) + 1):
        val = ws.cell(row=row, column=col).value
        if val is not None:
            data[headers[col-1]] = val
    return data

def is_bad_row(name_str):
    """Check if this row is a scored duplicate or broken"""
    if not name_str:
        return True
    name = str(name_str).strip()
    if re.match(r'^\[[\d.]+\]', name):
        return True
    if re.match(r'^\[AD:\d+\]', name):
        return True
    if 'LINE 48' in name:
        return True
    if name == '' or name == 'None':
        return True
    return False

def normalize_url(url):
    """Normalize YouTube URL for dedup"""
    if not url:
        return None
    url = str(url).strip().lower()
    url = url.replace('https://', '').replace('http://', '')
    url = url.rstrip('/')
    return url

def count_non_none(data, headers):
    """Count how many fields have data"""
    count = 0
    for h in headers:
        if h in data and data[h] is not None and str(data[h]).strip() != '':
            count += 1
    return count

def main():
    print("Loading workbook...")
    wb = openpyxl.load_workbook('YouTubers_Report_v2_FULL.xlsx')
    ws = wb.active

    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    print(f"Headers: {headers}")
    print(f"Original rows: {ws.max_row - 1}")

    # Step 1: Load all rows into dicts
    all_rows = []
    for row in range(2, ws.max_row + 1):
        data = load_row_data(ws, row, headers)
        name = data.get('Channel Name')
        if name and not is_bad_row(name):
            all_rows.append(data)

    print(f"After removing bad rows: {len(all_rows)}")

    # Step 2: Deduplicate by URL
    url_map = {}
    deduped = []
    for row_data in all_rows:
        url = normalize_url(row_data.get('YouTube URL'))
        if url is None:
            # Keep rows without URLs (might have other data)
            # But skip if no name either
            if row_data.get('Channel Name'):
                deduped.append(row_data)
            continue

        if url in url_map:
            # Keep the one with more data
            existing_idx = url_map[url]
            existing_count = count_non_none(deduped[existing_idx], headers)
            new_count = count_non_none(row_data, headers)
            if new_count > existing_count:
                deduped[existing_idx] = row_data
        else:
            url_map[url] = len(deduped)
            deduped.append(row_data)

    print(f"After URL dedup: {len(deduped)}")

    # Step 3: Also deduplicate by normalized channel name
    name_map = {}
    final_rows = []
    for row_data in deduped:
        name = str(row_data.get('Channel Name', '')).strip().lower()
        # Remove emojis and special chars for comparison
        clean_name = re.sub(r'[^\w\s]', '', name).strip()

        if clean_name in name_map:
            existing_idx = name_map[clean_name]
            existing_count = count_non_none(final_rows[existing_idx], headers)
            new_count = count_non_none(row_data, headers)
            if new_count > existing_count:
                # Merge: keep new but fill in any missing fields from old
                old = final_rows[existing_idx]
                for h in headers:
                    if (h not in row_data or row_data[h] is None) and h in old and old[h] is not None:
                        row_data[h] = old[h]
                final_rows[existing_idx] = row_data
        else:
            name_map[clean_name] = len(final_rows)
            final_rows.append(row_data)

    print(f"After name dedup: {len(final_rows)}")

    # Step 4: Enrich existing Uniminepool entries
    for row_data in final_rows:
        name = str(row_data.get('Channel Name', ''))

        # Enrich "Best Passive Income Programs Online (Main)"
        if 'Best Passive Income' in name:
            if not row_data.get('Other'):
                row_data['Other'] = 'linktr.ee/cryptojaime, Telegram channel: 12K subs'
            if not row_data.get('Notes') or 'Linktree' not in str(row_data.get('Notes', '')):
                existing_notes = str(row_data.get('Notes', ''))
                row_data['Notes'] = existing_notes + ' Linktree: linktr.ee/cryptojaime. Covers BSC yield farms, DeFi. Also has 2nd channel @bestpassiveincomeprogramso6200 (4.8K subs). TG channel @jaimesorianocrypto has 12K subs.'

        # Enrich CryptoLifestyle
        if name == 'CryptoLifestyle':
            if not row_data.get('Country') or row_data['Country'] == 'English':
                row_data['Country'] = 'USA / English'
            if not row_data.get('Notes') or 'Gruzewski' not in str(row_data.get('Notes', '')):
                existing_notes = str(row_data.get('Notes', ''))
                if 'P. Gruzewski' not in existing_notes:
                    row_data['Notes'] = existing_notes

        # Fix "Earn Online" - Tim Waider
        if name == 'Earn Online':
            if not row_data.get('Email'):
                row_data['Email'] = ''  # no verified email found
            if not row_data.get('Country') or row_data['Country'] == 'English':
                row_data['Country'] = 'English'

    # Step 5: Add new channels from research
    new_channels = [
        {
            'Channel Name': 'Best Passive Income Programs Online (2nd Channel)',
            'Subscribers': '4,800',
            'YouTube URL': 'https://youtube.com/@bestpassiveincomeprogramso6200',
            'Category': 'Uniminepool',
            'Country': 'USA / English',
            'Niche': 'Passive Income / Crypto',
            'Quality': 'LOW-MEDIUM',
            'Fake Risk': 'LOW',
            'Avg Views': '~200-500',
            'Upload Freq': 'Irregular',
            'Email': 'dailyresiduals@gmail.com',
            'Telegram': 't.me/jaimesorianocrypto',
            'Contact Quality': 'HIGH',
            'Role': 'Active Promoter',
            'Notes': '2nd channel of Jaime Soriano (main channel: 19.6K subs). Same contact info. Smaller reach but same audience niche.',
            'Verdict': 'MAYBE',
            'Verdict Reason': 'Secondary channel of established promoter. Very small reach at 4.8K subs, but same dedicated passive income audience as main channel.',
        },
        {
            'Channel Name': 'Mining Chamber',
            'Subscribers': '56,600',
            'YouTube URL': 'https://youtube.com/c/MiningChamber',
            'Category': 'Crypto Mining',
            'Country': 'English',
            'Niche': 'Crypto Mining / Hardware',
            'Quality': 'MEDIUM-HIGH',
            'Fake Risk': 'LOW',
            'Upload Freq': 'Regular',
            'Role': 'Educator / Reviewer',
            'Notes': 'Community for mining enthusiasts. Covers hardware, cloud mining, mining accessibility. Good potential for copy-trading as alt passive income.',
            'Verdict': 'MAYBE',
            'Verdict Reason': 'Legitimate mining community channel (56K subs). Audience interested in passive crypto income. Cloud mining coverage aligns with copy-trading pitch.',
        },
        {
            'Channel Name': 'Passive Crypto Mining',
            'Subscribers': '54,000',
            'YouTube URL': 'https://youtube.com/c/PassiveCryptoMining',
            'Category': 'Crypto Mining',
            'Country': 'English',
            'Niche': 'Passive Income / Mining',
            'Quality': 'MEDIUM',
            'Fake Risk': 'MEDIUM',
            'Upload Freq': 'Regular',
            'Role': 'Educator / Promoter',
            'Notes': 'Focused on earning passive income through crypto mining. Operating since Jan 2021. Covers various mining platforms and passive strategies.',
            'Verdict': 'MAYBE',
            'Verdict Reason': 'Good niche alignment (54K subs, passive income focus). Audience already interested in passive crypto earning methods. Worth manual review.',
        },
        {
            'Channel Name': 'HashRaptor',
            'Subscribers': '47,700',
            'YouTube URL': 'https://youtube.com/c/HashRaptor',
            'Category': 'Crypto Mining',
            'Country': 'USA',
            'Niche': 'Mining / Smart Home / Tech',
            'Quality': 'MEDIUM-HIGH',
            'Fake Risk': 'LOW',
            'Upload Freq': 'Regular',
            'Role': 'Educator / Reviewer',
            'Notes': 'Covers mining, cryptocurrency, home automation. Tech-savvy audience interested in passive setups.',
            'Verdict': 'MAYBE',
            'Verdict Reason': 'Tech-focused mining channel (47K subs) with genuine audience. Smart home + mining overlap suggests tech-savvy viewers open to automated trading.',
        },
        {
            'Channel Name': 'Mad Electron Engineering',
            'Subscribers': '38,900',
            'YouTube URL': 'https://youtube.com/c/MadElectronEngineering',
            'Category': 'Crypto Mining',
            'Country': 'USA',
            'Niche': 'Renewable Energy / Mining / Tech',
            'Quality': 'MEDIUM-HIGH',
            'Fake Risk': 'LOW',
            'Upload Freq': 'Regular',
            'Role': 'Educator',
            'Notes': 'Renewable energy + crypto mining experiments. Circuit design. Engineering-focused audience.',
            'Verdict': 'MAYBE',
            'Verdict Reason': 'Unique angle - renewable energy meets crypto mining (38K subs). Engineering audience may appreciate data-driven copy-trading approach.',
        },
        {
            'Channel Name': 'Cultivate Crypto',
            'Subscribers': '34,700',
            'YouTube URL': 'https://youtube.com/c/CultivateCrypto',
            'Category': 'Crypto Mining',
            'Country': 'Japan',
            'Niche': 'Crypto / Mining',
            'Quality': 'MEDIUM',
            'Fake Risk': 'LOW',
            'Upload Freq': 'Regular',
            'Role': 'Educator',
            'Notes': 'Japan-based crypto mining channel. Covers various mining topics.',
            'Verdict': 'MAYBE',
            'Verdict Reason': 'Japan-based mining channel (34K subs). Interesting geographic market (Japan crypto-friendly). Small but potentially engaged audience.',
        },
        {
            'Channel Name': 'ChumpChangeXD Mining & Crypto',
            'Subscribers': '30,000',
            'YouTube URL': 'https://youtube.com/c/ChumpChangeXD',
            'Category': 'Crypto Mining',
            'Country': 'English',
            'Niche': 'Mining / Crypto',
            'Quality': 'MEDIUM',
            'Fake Risk': 'LOW',
            'Upload Freq': 'Regular',
            'Role': 'Educator',
            'Notes': 'Mining and cryptocurrency focused content. Community-driven.',
            'Verdict': 'MAYBE',
            'Verdict Reason': 'Community-driven mining channel (30K subs). Mining audience overlaps with passive income seekers. Worth outreach if verified.',
        },
        {
            'Channel Name': 'Genesis Mining',
            'Subscribers': '26,400',
            'YouTube URL': 'https://youtube.com/c/GenesisMining',
            'Category': 'Cloud Mining',
            'Country': 'Germany',
            'Niche': 'Cloud Mining / Trading',
            'Quality': 'HIGH',
            'Fake Risk': 'LOW',
            'Upload Freq': 'Irregular',
            'Role': 'Platform (official)',
            'Notes': 'Official channel of Genesis Mining - major cloud mining provider. Also offers algorithmic trading. Brand-level contact.',
            'Verdict': 'MAYBE',
            'Verdict Reason': 'Official channel of major cloud mining brand (26K subs). Brand partnership potential rather than influencer outreach. High credibility.',
        },
        {
            'Channel Name': 'Crypto Mining Belgium',
            'Subscribers': '22,600',
            'YouTube URL': 'https://youtube.com/c/CryptoMiningBelgium',
            'Category': 'Crypto Mining',
            'Country': 'Belgium',
            'Niche': 'Mining / Passive Income',
            'Quality': 'MEDIUM-HIGH',
            'Fake Risk': 'LOW',
            'Avg Views': '~1,000-3,000',
            'Upload Freq': 'Regular',
            'Role': 'Educator',
            'Notes': '240 videos since 2016. Shows building long-term passive income with crypto. Journey to financial freedom. Experienced creator.',
            'Verdict': 'MAYBE',
            'Verdict Reason': 'Long-running Belgian mining channel (22K subs, since 2016, 240 videos). Passive income focus aligns well. Experienced and consistent creator.',
        },
        {
            'Channel Name': 'The Mining KiiNG',
            'Subscribers': '21,200',
            'YouTube URL': 'https://youtube.com/c/TheMiningKiiNG',
            'Category': 'Crypto Mining',
            'Country': 'USA',
            'Niche': 'GPU Mining / Education',
            'Quality': 'MEDIUM',
            'Fake Risk': 'LOW',
            'Upload Freq': 'Regular',
            'Role': 'Educator',
            'Notes': 'GPU mining fundamentals, overclocking, BIOS modding. Technical audience.',
            'Verdict': 'MAYBE',
            'Verdict Reason': 'Technical mining educator (21K subs). Audience invested in crypto mining infrastructure. May be open to reviewing trading platforms.',
        },
        {
            'Channel Name': 'Understanding Crypto - Herve',
            'Subscribers': '13,700',
            'YouTube URL': 'https://youtube.com/c/UnderstandingCryptoHerve',
            'Category': 'Crypto Education',
            'Country': 'UAE',
            'Niche': 'Crypto / Trading / Education',
            'Quality': 'MEDIUM',
            'Fake Risk': 'LOW',
            'Upload Freq': 'Regular',
            'Role': 'Educator',
            'Notes': 'Tutorials, crypto research and technical analysis. UAE-based.',
            'Verdict': 'MAYBE',
            'Verdict Reason': 'UAE-based crypto educator (13K subs). Technical analysis focus aligns with trading platform promotion. Small but engaged niche.',
        },
        {
            'Channel Name': 'DXB Crypto',
            'Subscribers': '10,800',
            'YouTube URL': 'https://youtube.com/c/DXBCrypto',
            'Category': 'Crypto',
            'Country': 'UAE',
            'Niche': 'Crypto / Community',
            'Quality': 'MEDIUM',
            'Fake Risk': 'LOW',
            'Upload Freq': 'Regular',
            'Role': 'Community',
            'Notes': 'Community-driven channel covering news, trends, and crypto projects. Dubai-based.',
            'Verdict': 'MAYBE',
            'Verdict Reason': 'Dubai-based crypto community channel (10K subs). UAE market is crypto-friendly. Community focus good for engagement.',
        },
    ]

    # Check that new channels don't duplicate existing ones
    existing_names_lower = set()
    existing_urls = set()
    for row_data in final_rows:
        name = str(row_data.get('Channel Name', '')).strip().lower()
        clean_name = re.sub(r'[^\w\s]', '', name).strip()
        existing_names_lower.add(clean_name)
        url = normalize_url(row_data.get('YouTube URL'))
        if url:
            existing_urls.add(url)

    added = 0
    for ch in new_channels:
        name = str(ch['Channel Name']).strip().lower()
        clean_name = re.sub(r'[^\w\s]', '', name).strip()
        url = normalize_url(ch.get('YouTube URL'))

        if clean_name in existing_names_lower:
            print(f"  SKIP (name dup): {ch['Channel Name']}")
            continue
        if url and url in existing_urls:
            print(f"  SKIP (URL dup): {ch['Channel Name']}")
            continue

        final_rows.append(ch)
        existing_names_lower.add(clean_name)
        if url:
            existing_urls.add(url)
        added += 1
        print(f"  ADDED: {ch['Channel Name']} ({ch['Subscribers']})")

    print(f"\nAdded {added} new channels")

    # Step 6: Sort - Rows with Verdict first (by subscriber count desc), then without verdict (by subs desc)
    def parse_subs(s):
        if not s:
            return 0
        s = str(s).replace(',', '').replace('+', '').strip()
        s = re.sub(r'[^\d]', '', s.split('.')[0] if '.' not in s else s)
        try:
            return int(s)
        except:
            return 0

    verdict_order = {'RECOMMEND': 0, 'MAYBE': 1, None: 2, '': 2}

    def sort_key(row_data):
        verdict = row_data.get('Verdict')
        v_order = verdict_order.get(verdict, 2)
        subs = parse_subs(row_data.get('Subscribers'))
        return (v_order, -subs)

    final_rows.sort(key=sort_key)

    # Step 7: Write to new workbook
    print(f"\nWriting {len(final_rows)} rows to cleaned workbook...")

    new_wb = openpyxl.Workbook()
    new_ws = new_wb.active
    new_ws.title = 'All Channels (Cleaned)'

    # Write headers
    header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for col, header in enumerate(headers, 1):
        cell = new_ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', wrap_text=True)
        cell.border = thin_border

    # Define fills for verdicts
    recommend_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
    maybe_fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
    no_verdict_fill = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')

    for idx, row_data in enumerate(final_rows):
        row_num = idx + 2

        # Set row number
        new_ws.cell(row=row_num, column=1, value=idx + 1)

        # Set all other data
        for col, header in enumerate(headers, 1):
            if header == '#':
                continue
            val = row_data.get(header)
            cell = new_ws.cell(row=row_num, column=col, value=val)
            cell.border = thin_border
            cell.alignment = Alignment(wrap_text=True, vertical='top')

            # Color by verdict
            verdict = row_data.get('Verdict')
            if verdict == 'RECOMMEND':
                cell.fill = recommend_fill
            elif verdict == 'MAYBE':
                cell.fill = maybe_fill
            elif not verdict:
                cell.fill = no_verdict_fill

    # Set column widths
    col_widths = {
        '#': 5, 'Verdict': 12, 'Channel Name': 35, 'Subscribers': 12,
        'YouTube URL': 45, 'Category': 18, 'Country': 18, 'Niche': 25,
        'Quality': 12, 'Fake Risk': 10, 'Avg Views': 15, 'View/Sub %': 10,
        'Upload Freq': 12, 'Community (1-10)': 10, 'Engagement': 12,
        'Email': 30, 'Telegram': 30, 'Instagram': 25, 'Twitter/X': 25,
        'WhatsApp': 15, 'Other': 35, 'Contact Quality': 12, 'Role': 20,
        'Notes': 50, 'Verdict Reason': 60
    }

    for col, header in enumerate(headers, 1):
        width = col_widths.get(header, 15)
        new_ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width

    # Freeze top row
    new_ws.freeze_panes = 'A2'

    # Auto filter
    new_ws.auto_filter.ref = f'A1:{openpyxl.utils.get_column_letter(len(headers))}{len(final_rows) + 1}'

    output_file = 'YouTubers_Report_v2_CLEANED.xlsx'
    new_wb.save(output_file)
    print(f"\nSaved to {output_file}")

    # Stats
    recommend_count = sum(1 for r in final_rows if r.get('Verdict') == 'RECOMMEND')
    maybe_count = sum(1 for r in final_rows if r.get('Verdict') == 'MAYBE')
    no_verdict = sum(1 for r in final_rows if not r.get('Verdict'))
    uniminepool_count = sum(1 for r in final_rows if 'unimine' in str(r.get('Category', '')).lower())

    print(f"\n=== FINAL STATS ===")
    print(f"Total channels: {len(final_rows)}")
    print(f"RECOMMEND: {recommend_count}")
    print(f"MAYBE: {maybe_count}")
    print(f"No verdict: {no_verdict}")
    print(f"Uniminepool-related: {uniminepool_count}")

if __name__ == '__main__':
    main()
