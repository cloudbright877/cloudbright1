"""Fix data integrity issues and apply final contacts."""
import openpyxl
import re
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Fix specific channels with wrong data + add new contacts
FIXES = {
    'HashRaptor': {
        'Email': 'HashRaptor@gmail.com',
        'Telegram': '',  # Remove wrong telegram
        'Twitter/X': 'x.com/HashRaptor',
        'Other': 'hashraptor.com, kit.co/HashRaptor',
        'Contact Quality': 'HIGH',
        'Country': 'USA',
        'Notes': '48.6K subs. Crypto mining, home automation, smart home, gaming.',
    },
    'Crypto Miner Tips': {
        'Instagram': 'instagram.com/cryptominertips',
        'Other': 'cryptominertips.com, Discord: discord.gg/2WbPZf2HRv, Facebook Group',
        'Contact Quality': 'HIGH',
    },
    'Money Maker - Avi Lev': {
        'Country': 'Israel',
    },
    'Passive Crypto Mining': {
        'Email': 'info.passivecrypto@gmail.com',
        'Country': 'USA',
    },
}

NEW_CONTACTS = {
    'Across The Rubicon': {
        'Telegram': 't.me/Rubicontwilly',
        'Twitter/X': 'x.com/acrosstherubico',
        'Other': 'patreon.com/rubiconresearch, Discord: discord.com/invite/rubicon-752590582274326680, rubiconinnercircle.com',
        'Contact Quality': 'MEDIUM',
        'Country': 'Australia',
        'Notes': 'Run by Benji (Gold Coast) & Tristan (Twilly). Web3, crypto cheat sheets, token breakdowns. 150K subs.',
    },
    'Full Value Dan': {
        'Twitter/X': 'x.com/fullvaluedan',
        'Instagram': 'instagram.com/fullvaluedan',
        'Notes': 'Real name Daniel Arreola, LA area. TikTok: @fullvaluedan, LinkedIn. Also on Eddy LIVE podcast.',
    },
    'Money Maker - Avi Lev': {
        'Instagram': 'instagram.com/moneymakeravllev',
        'Notes': 'Financial education. Stocks, dividends, crypto. 179K subs. TikTok: @avilev_mm. CEO Dript IV Therapy.',
    },
    'Your Friend Andy': {
        'Instagram': 'instagram.com/your.friend.andy',
        'Other': 'yourfriendandy.com, yourfriendandy.bio, muskminers.com (partner), Podcast: Andy and Friends, Newsletter: email.yourfriendandy.com',
    },
    'NEXT GEN CRYPTO': {
        'Twitter/X': 'x.com/NextGenCrypto',
        'Instagram': 'instagram.com/next.gen.crypto',
        'Other': 'nextgencrypto.online, patreon.com/NextGenCryptoYT',
        'Contact Quality': 'MEDIUM',
        'Notes': '1.52M subs. Major crypto channel.',
    },
    'Crypto Peak': {
        'Instagram': 'instagram.com/cryptopeakyt',
        'Contact Quality': 'MEDIUM',
        'Notes': '272K subs. Crypto analysis.',
    },
    'CryptoDexWorld': {
        'Twitter/X': 'x.com/cryptodexworld',
        'Contact Quality': 'MEDIUM',
        'Notes': '134K subs. Crypto education.',
    },
    'Crypto Bolt': {
        'Twitter/X': 'x.com/CryptoBoltx',
        'Other': 'Fiverr: ammar76358 (paid promo $70)',
        'Contact Quality': 'MEDIUM',
        'Notes': 'Ammar Khan, digital marketer. 114K subs. Offers paid crypto/NFT promo on YT.',
    },
    'P4 Provider': {
        'Email': 'contact@p4provider.com',
        'Twitter/X': 'x.com/P4Provider',
        'Notes': 'Hafiz Muhammad Tanveer (CEO). Shop #60, 3rd Floor, Misaq-ul-Mall, Faisalabad. Binance KOL @binancepakistani.',
    },
    'Trade With Aamir': {
        'Telegram': 't.me/TradeWithAamir',
        'Contact Quality': 'MEDIUM',
        'Notes': '188K subs. Associated with TreasureNFT platform.',
    },
    'MANU CHHINA': {
        'Instagram': 'instagram.com/digitalmanuchhina',
        'Twitter/X': 'x.com/manaman_chhina',
        'Other': 'digitalmanuchhina.in',
        'Contact Quality': 'MEDIUM',
        'Country': 'India',
        'Notes': '101K subs. Digital marketing courses, stock market, crypto trading.',
    },
    'Albarizone': {
        'Telegram': 't.me/albarizonpk',
        'Other': 'linktr.ee/albarizon.pk',
        'Contact Quality': 'MEDIUM',
        'Notes': '70.1K subs. Pi Network, crypto airdrops, coin listings.',
    },
    'MrCrypto PH': {
        'Telegram': 't.me/MrCryptoPHAnnouncement',
        'Twitter/X': 'x.com/MrCryptoPH1',
        'Contact Quality': 'MEDIUM',
        'Country': 'Philippines',
        'Notes': '35.3K subs. Philippines crypto channel.',
    },
    'Crypto Scholar': {
        'Twitter/X': 'x.com/icryptoscholar',
        'Other': 'cryptoscholar.io, cryptoscholar.io/contact-crypto-scholar',
        'Contact Quality': 'MEDIUM',
        'Notes': '31.7K subs. Educational crypto content. Web3/Lens: cryptoscholar.lens.',
    },
    'Crypto Empire': {
        'Twitter/X': 'x.com/csellcrypto',
        'Other': 'cryptoempiremastermind.com',
    },
    'Genesis Mining': {
        'Telegram': 't.me/genesismining',
        'Instagram': 'instagram.com/genesismining_official',
        'Country': 'Hong Kong',
        'Notes': 'Cloud mining company, HQ Hong Kong, ops in Iceland. Not onboarding new since March 2020. Beware fakes.',
    },
    'ScamFinder': {
        'Twitter/X': 'x.com/ScamFinder13',
        'Other': 'scamfinder.net',
        'Contact Quality': 'MEDIUM',
        'Notes': 'Publishes crypto scam reviews. BitHarvest review site.',
    },
    'Mining Chamber': {
        'Other': 'miningchamber.com, linktr.ee/miningchamber, Discord: discord.com/invite/mining-chamber-621000102403637269 (5.5K members)',
    },
    'Passive Crypto Mining': {
        'Email': 'info.passivecrypto@gmail.com',
        'Country': 'USA',
    },
    'HashRaptor': {
        'Instagram': 'instagram.com/hashraptor',
    },
    'CryptoTab Browser': {
        'Twitter/X': 'x.com/CryptoTabnet',
        'Instagram': 'instagram.com/cryptotabbrowser',
    },
    'Bank Of Jharol': {
        'Instagram': 'instagram.com/therealj_smith',
        'Other': 'tiktok.com/@bankofjharol, skool.com/@jharol-smith-2234 (113 members)',
        'Contact Quality': 'MEDIUM',
        'Country': 'USA',
        'Notes': 'Jharol Smith, DMV area. Crypto platforms, affiliate marketing. 21.7K subs.',
    },
    'AmG Arewa Tv': {
        'Country': 'Nigeria',
    },
    'Bangla Tech-iELA': {
        'Country': 'Bangladesh',
    },
}


def is_bad_row(name_str):
    if not name_str:
        return True
    name = str(name_str).strip()
    if re.match(r'^\[[\d.]+\]', name):
        return True
    if re.match(r'^\[AD:\d+\]', name):
        return True
    if 'LINE 48' in name:
        return True
    if name.startswith('---'):
        return True
    return False


def normalize_url(url):
    if not url:
        return None
    return str(url).strip().lower().replace('https://', '').replace('http://', '').rstrip('/')


def parse_subs(s):
    if not s:
        return 0
    s = str(s).replace(',', '').replace('+', '').strip()
    s = re.sub(r'[^\d]', '', s.split('.')[0] if '.' not in s else s)
    try:
        return int(s)
    except:
        return 0


def count_fields(data, headers):
    return sum(1 for h in headers if data.get(h) and str(data[h]).strip())


def main():
    wb = openpyxl.load_workbook('YouTubers_Report_v2_CLEANED.xlsx')
    ws = wb.active
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]

    all_rows = []
    for row in range(2, ws.max_row + 1):
        data = {}
        for col in range(1, len(headers) + 1):
            val = ws.cell(row=row, column=col).value
            if val is not None:
                data[headers[col - 1]] = val
        name = data.get('Channel Name')
        if name and not is_bad_row(name):
            all_rows.append(data)

    # Dedup by URL
    url_map = {}
    deduped = []
    for rd in all_rows:
        url = normalize_url(rd.get('YouTube URL'))
        if url is None:
            if rd.get('Channel Name'):
                deduped.append(rd)
            continue
        if url in url_map:
            idx = url_map[url]
            if count_fields(rd, headers) > count_fields(deduped[idx], headers):
                deduped[idx] = rd
        else:
            url_map[url] = len(deduped)
            deduped.append(rd)

    # Dedup by name
    name_map = {}
    final = []
    for rd in deduped:
        nm = str(rd.get('Channel Name', '')).strip().lower()
        clean = re.sub(r'[^\w\s]', '', nm).strip()
        if clean in name_map:
            idx = name_map[clean]
            if count_fields(rd, headers) > count_fields(final[idx], headers):
                old = final[idx]
                for h in headers:
                    if (h not in rd or rd[h] is None) and h in old and old[h] is not None:
                        rd[h] = old[h]
                final[idx] = rd
        else:
            name_map[clean] = len(final)
            final.append(rd)

    print(f"Channels: {len(final)}")

    # Apply fixes (force overwrite)
    fixed = 0
    for rd in final:
        name = str(rd.get('Channel Name', '')).strip()
        if name in FIXES:
            for field, value in FIXES[name].items():
                rd[field] = value if value else None
                fixed += 1
    print(f"Fixes applied: {fixed}")

    # Apply new contacts (only if empty)
    updated = 0
    for rd in final:
        name = str(rd.get('Channel Name', '')).strip()
        if name in NEW_CONTACTS:
            for field, value in NEW_CONTACTS[name].items():
                current = rd.get(field)
                if not current or str(current).strip() == '':
                    rd[field] = value
                    updated += 1
    print(f"New contact updates: {updated}")

    # Sort by category
    verdict_order = {'RECOMMEND': 0, 'MAYBE': 1}

    def sort_key(rd):
        cat = rd.get('Category') or 'Uncategorized'
        cat = str(cat).strip()
        if cat in ('None', ''):
            cat = 'zzz_Uncategorized'
        v = verdict_order.get(rd.get('Verdict'), 2)
        subs = parse_subs(rd.get('Subscribers'))
        return (cat.lower(), v, -subs)

    final.sort(key=sort_key)

    # Write new workbook
    new_wb = openpyxl.Workbook()
    nws = new_wb.active
    nws.title = 'Channels by Project'

    hdr_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
    hdr_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    rec_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
    may_fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
    nv_fill = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
    cat_fill = PatternFill(start_color='D6DCE4', end_color='D6DCE4', fill_type='solid')
    cat_font = Font(name='Calibri', size=11, bold=True, color='1F4E79')
    bdr = Border(left=Side(style='thin'), right=Side(style='thin'),
                 top=Side(style='thin'), bottom=Side(style='thin'))

    for col, h in enumerate(headers, 1):
        c = nws.cell(row=1, column=col, value=h)
        c.fill = hdr_fill
        c.font = hdr_font
        c.alignment = Alignment(horizontal='center', wrap_text=True)
        c.border = bdr

    cur = 2
    cur_cat = None
    num = 1

    for rd in final:
        cat = rd.get('Category') or 'Uncategorized'
        cat_d = str(cat).strip()
        if cat_d in ('None', ''):
            cat_d = 'Uncategorized'

        if cat_d != cur_cat:
            cur_cat = cat_d
            for col in range(1, len(headers) + 1):
                c = nws.cell(row=cur, column=col)
                c.fill = cat_fill
                c.border = bdr
            nws.cell(row=cur, column=3, value=f"--- {cat_d} ---").font = cat_font
            nws.cell(row=cur, column=3).fill = cat_fill
            cur += 1

        nws.cell(row=cur, column=1, value=num).border = bdr
        for col, h in enumerate(headers, 1):
            if h == '#':
                continue
            val = rd.get(h)
            c = nws.cell(row=cur, column=col, value=val)
            c.border = bdr
            c.alignment = Alignment(wrap_text=True, vertical='top')
            v = rd.get('Verdict')
            if v == 'RECOMMEND':
                c.fill = rec_fill
            elif v == 'MAYBE':
                c.fill = may_fill
            elif not v:
                c.fill = nv_fill

        num += 1
        cur += 1

    widths = {'#': 5, 'Verdict': 12, 'Channel Name': 35, 'Subscribers': 12,
              'YouTube URL': 45, 'Category': 18, 'Country': 18, 'Niche': 25,
              'Quality': 12, 'Fake Risk': 10, 'Avg Views': 15, 'View/Sub %': 10,
              'Upload Freq': 12, 'Community (1-10)': 10, 'Engagement': 12,
              'Email': 30, 'Telegram': 30, 'Instagram': 25, 'Twitter/X': 25,
              'WhatsApp': 15, 'Other': 35, 'Contact Quality': 12, 'Role': 20,
              'Notes': 50, 'Verdict Reason': 60}
    for col, h in enumerate(headers, 1):
        nws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = widths.get(h, 15)

    nws.freeze_panes = 'A2'
    nws.auto_filter.ref = f'A1:{openpyxl.utils.get_column_letter(len(headers))}{cur - 1}'

    new_wb.save('YouTubers_Report_v2_CLEANED.xlsx')

    hc = sum(1 for rd in final if any(rd.get(f) and str(rd[f]).strip() for f in ['Email', 'Telegram', 'WhatsApp']))
    twit = sum(1 for rd in final if rd.get('Twitter/X') and str(rd['Twitter/X']).strip())
    ig = sum(1 for rd in final if rd.get('Instagram') and str(rd['Instagram']).strip())
    print(f"\nWith direct contacts (E/TG/WA): {hc}/{len(final)}")
    print(f"With Twitter/X: {twit}")
    print(f"With Instagram: {ig}")
    print(f"Total with ANY social: {sum(1 for rd in final if any(rd.get(f) and str(rd[f]).strip() for f in ['Email', 'Telegram', 'WhatsApp', 'Twitter/X', 'Instagram']))}")
    print(f"RECOMMEND: {sum(1 for r in final if r.get('Verdict')=='RECOMMEND')}")
    print(f"MAYBE: {sum(1 for r in final if r.get('Verdict')=='MAYBE')}")
    print("Done!")


if __name__ == '__main__':
    main()
