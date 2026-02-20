"""
Update contacts for channels and sort by Category (project).
Creates a brand new workbook to avoid cell clearing issues.
"""
import openpyxl
import re
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

CONTACT_UPDATES = {
    'Tabraiz Shams': {
        'Email': 'tabraizvlogs@gmail.com',
        'Telegram': 't.me/ShamsTabraiz',
        'Instagram': 'instagram.com/shams__tabraiz',
        'Contact Quality': 'HIGH',
    },
    'P4 Provider': {
        'Telegram': 't.me/P4Provider, DM: @HMTanveer',
        'WhatsApp': '0328-8855333',
        'Other': 'p4provider.com, linktr.ee/p4provider',
        'Contact Quality': 'HIGH',
        'Notes': 'Hafiz Muhammad Tanveer (CEO). Pakistan trading institute. Binance KOL.',
    },
    'The Hobbyist Miner': {
        'Other': 'thehobbyistminer.io, Discord community',
        'Country': 'USA',
        'Contact Quality': 'MEDIUM',
    },
    'IITian in Crypto': {
        'Twitter/X': 'x.com/IITian_InCrypto',
        'Country': 'India',
    },
    'Christian Rauchenwald': {
        'Other': 'christianrauchenwald.com',
        'Contact Quality': 'MEDIUM',
    },
    'Blockchain SAGE': {
        'Telegram': 't.me/Sageblockchain',
        'Twitter/X': 'x.com/Bl0ckchainsage',
        'Contact Quality': 'MEDIUM',
    },
    'Crypto Empire': {
        'Telegram': 't.me/cryptoempireee, DM: @csell34',
        'Other': 'cryptoempireco.io, Discord',
        'Contact Quality': 'MEDIUM',
    },
    'Crypto Bee': {
        'Email': 'cryptobeeyt@gmail.com',
        'Telegram': 't.me/officialcryptobee, DM: @Jack_Berdul',
        'Contact Quality': 'HIGH',
    },
    'Darryl Boo': {
        'Telegram': 'DM: @cryptoboycash',
        'Twitter/X': 'x.com/darryl_boo',
        'Other': 'linktr.ee/darrylboo, TikTok: @darrylboocryptoboy',
        'Contact Quality': 'MEDIUM',
    },
    'Bitbull': {
        'Other': 'bitbullpro.net',
        'Country': 'Germany',
    },
    'Full Value Dan': {
        'Other': 'danteachescrypto.com, Telegram membership bot',
        'Contact Quality': 'MEDIUM',
    },
    'Your Friend Andy': {
        'Other': 'Partner: muskminers.com',
        'Country': 'USA',
    },
}


def parse_subs(s):
    if not s:
        return 0
    s = str(s).replace(',', '').replace('+', '').strip()
    s = re.sub(r'[^\d]', '', s.split('.')[0] if '.' not in s else s)
    try:
        return int(s)
    except:
        return 0


def main():
    print("Loading workbook...")
    wb = openpyxl.load_workbook('YouTubers_Report_v2_CLEANED.xlsx')
    ws = wb.active

    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    header_idx = {h: i + 1 for i, h in enumerate(headers)}
    total_rows = ws.max_row - 1
    print(f"Total channels: {total_rows}")

    # Step 1: Read all data into list
    all_rows = []
    for row in range(2, ws.max_row + 1):
        row_data = {}
        for col in range(1, len(headers) + 1):
            val = ws.cell(row=row, column=col).value
            if val is not None:
                row_data[headers[col - 1]] = val
        if row_data.get('Channel Name'):
            all_rows.append(row_data)

    print(f"Read {len(all_rows)} channels")

    # Step 2: Apply contact updates to data dicts
    updated = 0
    for row_data in all_rows:
        name = str(row_data.get('Channel Name', '')).strip()
        if name in CONTACT_UPDATES:
            for field, value in CONTACT_UPDATES[name].items():
                current = row_data.get(field)
                if not current or str(current).strip() == '':
                    row_data[field] = value
                    updated += 1

    print(f"Applied {updated} contact updates")

    # Step 3: Sort by Category, then Verdict, then Subscribers
    verdict_order = {'RECOMMEND': 0, 'MAYBE': 1}

    def norm_cat(cat):
        if not cat or str(cat).strip() in ('None', ''):
            return 'zzz_Uncategorized'
        return str(cat).strip()

    def sort_key(rd):
        cat = norm_cat(rd.get('Category'))
        verdict = rd.get('Verdict')
        v_order = verdict_order.get(verdict, 2)
        subs = parse_subs(rd.get('Subscribers'))
        return (cat.lower(), v_order, -subs)

    all_rows.sort(key=sort_key)

    # Step 4: Create brand new workbook
    new_wb = openpyxl.Workbook()
    new_ws = new_wb.active
    new_ws.title = 'Channels by Project'

    # Styles
    header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    recommend_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
    maybe_fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
    no_verdict_fill = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
    cat_fill = PatternFill(start_color='D6DCE4', end_color='D6DCE4', fill_type='solid')
    cat_font = Font(name='Calibri', size=11, bold=True, color='1F4E79')
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    # Write headers
    for col, h in enumerate(headers, 1):
        cell = new_ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', wrap_text=True)
        cell.border = thin_border

    # Write data with category separators
    current_row = 2
    current_cat = None
    num = 1

    for row_data in all_rows:
        cat = row_data.get('Category') or 'Uncategorized'
        cat_display = str(cat).strip()
        if cat_display in ('None', ''):
            cat_display = 'Uncategorized'

        # Category separator
        if cat_display != current_cat:
            current_cat = cat_display
            for col in range(1, len(headers) + 1):
                c = new_ws.cell(row=current_row, column=col)
                c.fill = cat_fill
                c.border = thin_border
            new_ws.cell(row=current_row, column=3, value=f"--- {cat_display} ---")
            new_ws.cell(row=current_row, column=3).font = cat_font
            new_ws.cell(row=current_row, column=3).fill = cat_fill
            current_row += 1

        # Data row
        new_ws.cell(row=current_row, column=1, value=num).border = thin_border
        for col, h in enumerate(headers, 1):
            if h == '#':
                continue
            val = row_data.get(h)
            cell = new_ws.cell(row=current_row, column=col, value=val)
            cell.border = thin_border
            cell.alignment = Alignment(wrap_text=True, vertical='top')

            verdict = row_data.get('Verdict')
            if verdict == 'RECOMMEND':
                cell.fill = recommend_fill
            elif verdict == 'MAYBE':
                cell.fill = maybe_fill
            elif not verdict:
                cell.fill = no_verdict_fill

        num += 1
        current_row += 1

    # Column widths
    col_widths = {
        '#': 5, 'Verdict': 12, 'Channel Name': 35, 'Subscribers': 12,
        'YouTube URL': 45, 'Category': 18, 'Country': 18, 'Niche': 25,
        'Quality': 12, 'Fake Risk': 10, 'Avg Views': 15, 'View/Sub %': 10,
        'Upload Freq': 12, 'Community (1-10)': 10, 'Engagement': 12,
        'Email': 30, 'Telegram': 30, 'Instagram': 25, 'Twitter/X': 25,
        'WhatsApp': 15, 'Other': 35, 'Contact Quality': 12, 'Role': 20,
        'Notes': 50, 'Verdict Reason': 60
    }
    for col, h in enumerate(headers, 1):
        w = col_widths.get(h, 15)
        new_ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = w

    new_ws.freeze_panes = 'A2'
    new_ws.auto_filter.ref = f'A1:{openpyxl.utils.get_column_letter(len(headers))}{current_row - 1}'

    output = 'YouTubers_Report_v2_CLEANED.xlsx'
    new_wb.save(output)
    print(f"\nSaved to {output}")

    # Stats
    cat_counts = {}
    for rd in all_rows:
        c = rd.get('Category') or 'Uncategorized'
        cs = str(c).strip()
        if cs in ('None', ''):
            cs = 'Uncategorized'
        cat_counts[cs] = cat_counts.get(cs, 0) + 1

    print(f"\n=== {len(cat_counts)} CATEGORIES ===")
    for c, n in sorted(cat_counts.items(), key=lambda x: -x[1]):
        print(f"  {c}: {n}")

    has_contact = 0
    for rd in all_rows:
        if any(rd.get(f) and str(rd.get(f)).strip() for f in ['Email', 'Telegram', 'WhatsApp']):
            has_contact += 1
    print(f"\nWith contacts: {has_contact}/{len(all_rows)}")
    print(f"Total rows written: {current_row - 2} (inc. {len(cat_counts)} separators)")


if __name__ == '__main__':
    main()
