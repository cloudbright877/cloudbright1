"""
Scan YouTube channels without Telegram contacts.
Extract channel description + last 5 video descriptions.
Search for t.me/ and telegram links.
"""
import subprocess
import json
import re
import sys
import time
import openpyxl

def get_channel_telegram(channel_url, channel_name):
    """Extract telegram links from channel description and recent videos."""
    telegrams = set()

    # Normalize URL
    url = channel_url.strip()
    if not url.startswith('http'):
        url = 'https://' + url

    # Try /videos endpoint
    videos_url = url.rstrip('/')
    if '/channel/' in videos_url or '/@' in videos_url:
        if not videos_url.endswith('/videos'):
            videos_url = videos_url.rstrip('/') + '/videos'

    try:
        # Get channel info + 5 recent videos
        result = subprocess.run(
            ['python', '-m', 'yt_dlp', '--dump-single-json', '--flat-playlist',
             '--playlist-items', '1-5', videos_url],
            capture_output=True, text=True, timeout=45,
            encoding='utf-8', errors='replace'
        )

        if result.stdout:
            data = json.loads(result.stdout)

            # Check channel description
            desc = data.get('description', '') or ''
            tg_links = re.findall(r'(?:https?://)?(?:t\.me|telegram\.me)/[A-Za-z0-9_]+', desc)
            for link in tg_links:
                if not link.startswith('http'):
                    link = 'https://' + link
                telegrams.add(link.replace('https://telegram.me/', 'https://t.me/').replace('https://t.me/', 't.me/'))

            # Check each video description
            entries = data.get('entries', [])
            video_ids = [e.get('id') for e in entries if e.get('id')]

            for vid_id in video_ids[:5]:
                try:
                    vresult = subprocess.run(
                        ['python', '-m', 'yt_dlp', '--dump-json', '--skip-download',
                         f'https://www.youtube.com/watch?v={vid_id}'],
                        capture_output=True, text=True, timeout=30,
                        encoding='utf-8', errors='replace'
                    )
                    if vresult.stdout:
                        vdata = json.loads(vresult.stdout)
                        vdesc = vdata.get('description', '') or ''
                        vtg = re.findall(r'(?:https?://)?(?:t\.me|telegram\.me)/[A-Za-z0-9_]+', vdesc)
                        for link in vtg:
                            if not link.startswith('http'):
                                link = 'https://' + link
                            telegrams.add(link.replace('https://telegram.me/', 't.me/').replace('https://t.me/', 't.me/').replace('http://t.me/', 't.me/'))
                except Exception:
                    pass

    except subprocess.TimeoutExpired:
        pass
    except Exception as e:
        pass

    # Filter out generic/bot links
    filtered = set()
    skip_patterns = ['t.me/share', 't.me/joinchat', 't.me/addstickers', 't.me/proxy', 't.me/socks']
    for tg in telegrams:
        clean = tg.lower().strip()
        if any(p in clean for p in skip_patterns):
            continue
        # Extract just the handle
        match = re.search(r't\.me/([A-Za-z0-9_]+)', tg)
        if match:
            handle = match.group(1)
            if len(handle) >= 3 and handle not in ('bot', 'share', 'proxy'):
                filtered.add(f't.me/{handle}')

    return filtered


def main():
    print("Loading spreadsheet...")
    wb = openpyxl.load_workbook('YouTubers_Report_v2_CLEANED.xlsx')
    ws = wb.active
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    col_map = {h: i+1 for i, h in enumerate(headers)}

    # Find channels without telegram
    targets = []
    for row in range(2, ws.max_row + 1):
        name = ws.cell(row=row, column=col_map['Channel Name']).value
        if not name or str(name).startswith('---'):
            continue
        tg = ws.cell(row=row, column=col_map['Telegram']).value
        if tg and str(tg).strip() not in ('', '-', 'None'):
            continue
        url = ws.cell(row=row, column=col_map['YouTube URL']).value
        if not url:
            continue

        subs_raw = ws.cell(row=row, column=col_map['Subscribers']).value
        s = str(subs_raw).replace(',', '').replace('+', '').strip() if subs_raw else '0'
        s = re.sub(r'[^\d]', '', s.split('.')[0] if '.' not in s else s)
        try:
            subs = int(s)
        except:
            subs = 0

        if subs >= 10000:
            targets.append((row, str(name).strip(), str(url).strip(), subs))

    targets.sort(key=lambda x: -x[3])  # Sort by subs desc
    print(f"Channels to scan: {len(targets)}")

    found_count = 0
    results = {}

    for i, (row, name, url, subs) in enumerate(targets):
        safe_name = ''.join(c if ord(c) < 128 else '?' for c in name)
        print(f"[{i+1}/{len(targets)}] {safe_name} ({subs:,} subs)...", end=' ', flush=True)

        tg_links = get_channel_telegram(url, name)

        if tg_links:
            tg_str = '; '.join(sorted(tg_links))
            print(f"FOUND: {tg_str}")
            results[row] = tg_str
            found_count += 1

            # Update spreadsheet
            ws.cell(row=row, column=col_map['Telegram'], value=tg_str)
        else:
            print("none")

        # Small delay to avoid rate limiting
        if (i + 1) % 10 == 0:
            time.sleep(1)

        # Save every 25 channels
        if (i + 1) % 25 == 0:
            wb.save('YouTubers_Report_v2_CLEANED.xlsx')
            print(f"  --- Saved. Found so far: {found_count}/{i+1} ---")

    # Final save
    wb.save('YouTubers_Report_v2_CLEANED.xlsx')

    print(f"\n{'='*50}")
    print(f"RESULTS: Found Telegram for {found_count}/{len(targets)} channels")
    print(f"{'='*50}")
    for row, tg in sorted(results.items()):
        name = ws.cell(row=row, column=col_map['Channel Name']).value
        print(f"  {name}: {tg}")


if __name__ == '__main__':
    main()
