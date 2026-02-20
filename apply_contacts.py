"""
Apply batch 2 contacts + re-sort.
"""
import openpyxl
import re
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

CONTACT_UPDATES = {
    'Tabraiz Shams': {
        'Email': 'tabraizvlogs@gmail.com',
        'Telegram': 't.me/ShamsTabraiz, DM promo: @Satoshi21m',
        'Twitter/X': 'x.com/janTabraiz',
        'Instagram': 'instagram.com/shams__tabraiz',
        'Contact Quality': 'HIGH',
        'Country': 'Pakistan',
        'Notes': 'Based in Pakistan (not India). Crypto airdrops/blockchain projects.',
    },
    'The Hobbyist Miner': {
        'Email': 'TheHobbyistMiner@gmail.com',
        'Twitter/X': 'x.com/HobbyistMiner',
        'Other': 'thehobbyistminer.io, Discord: discord.com/invite/bRqFMBxTrv',
        'Country': 'USA',
        'Contact Quality': 'HIGH',
    },
    'IITian in Crypto': {
        'Email': 'team.iitianincrypto@gmail.com',
        'Telegram': 't.me/IITian_incrypto',
        'Twitter/X': 'x.com/IITian_InCrypto',
        'Instagram': 'instagram.com/iitian.incrypto',
        'Contact Quality': 'HIGH',
        'Country': 'India',
        'Notes': 'Umang Gupta, IIT Madras grad, Pune. 500K+ followers. IDO/IGO strategies.',
    },
    'Christian Rauchenwald': {
        'Email': 'christian@convertain.com',
        'Twitter/X': 'x.com/ChristianRauche',
        'Instagram': 'instagram.com/christianrauchenwald',
        'Other': 'christianrauchenwald.com, LinkedIn, GitHub: RauchenwaldC, CEO Convertain Holdings GmbH, HK company',
        'Contact Quality': 'HIGH',
    },
    'Token Galaxy': {
        'Telegram': 't.me/Nexora_Global (project)',
        'Other': 'nexoracrypto.com, CoinRise agency: info@coinrise.io',
        'Contact Quality': 'MEDIUM',
    },
    'Darryl Boo': {
        'Telegram': 'DM: @cryptoboycash',
        'Twitter/X': 'x.com/darryl_boo',
        'Other': 'linktr.ee/darrylboo, TikTok: @darrylboocryptoboy',
        'Contact Quality': 'MEDIUM',
    },
    'Mr. Tech Official': {
        'Telegram': 't.me/rahulint (40.8K subs)',
        'Other': 'mrtechofficial.com',
        'Country': 'India',
        'Contact Quality': 'MEDIUM',
    },
    'Blockchain SAGE': {
        'Telegram': 't.me/Sageblockchain',
        'Twitter/X': 'x.com/Bl0ckchainsage',
        'Contact Quality': 'MEDIUM',
    },
    'Crypto Empire': {
        'Telegram': 't.me/cryptoempireee, DM: @csell34',
        'Other': 'cryptoempireco.io, Discord',
        'Contact Quality': 'MEDIUM',
    },
    'Crypto Bee': {
        'Email': 'cryptobeeyt@gmail.com',
        'Telegram': 't.me/officialcryptobee, DM: @Jack_Berdul',
        'Contact Quality': 'HIGH',
    },
    'P4 Provider': {
        'Telegram': 't.me/P4Provider, DM: @HMTanveer',
        'WhatsApp': '0328-8855333',
        'Other': 'p4provider.com, linktr.ee/p4provider',
        'Contact Quality': 'HIGH',
    },
    'Full Value Dan': {
        'Other': 'danteachescrypto.com, Telegram membership bot',
        'Contact Quality': 'MEDIUM',
    },
    'Your Friend Andy': {
        'Other': 'Partner: muskminers.com',
        'Country': 'USA',
    },
    'Bitbull': {
        'Other': 'bitbullpro.net',
        'Country': 'Germany',
    },
    'Best Passive Income Programs Online (Main)': {
        'Other': 'linktr.ee/cryptojaime, TG channel: 12K subs',
    },
}


def is_bad_row(name_str):
    if not name_str:
        return True
    name = str(name_str).strip()
    if re.match(r'^\[[\d.]+\]', name):
        return True
    if re.match(r'^\[AD:\d+\]', name):
        return True
    if 'LINE 48' in name:
        return True
    if name.startswith('---'):
        return True
    return False


def normalize_url(url):
    if not url:
        return None
    return str(url).strip().lower().replace('https://', '').replace('http://', '').rstrip('/')


def parse_subs(s):
    if not s:
        return 0
    s = str(s).replace(',', '').replace('+', '').strip()
    s = re.sub(r'[^\d]', '', s.split('.')[0] if '.' not in s else s)
    try:
        return int(s)
    except:
        return 0


def count_fields(data, headers):
    return sum(1 for h in headers if data.get(h) and str(data[h]).strip())


def main():
    wb = openpyxl.load_workbook('YouTubers_Report_v2_CLEANED.xlsx')
    ws = wb.active
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]

    # Read, filter bad rows
    all_rows = []
    for row in range(2, ws.max_row + 1):
        data = {}
        for col in range(1, len(headers) + 1):
            val = ws.cell(row=row, column=col).value
            if val is not None:
                data[headers[col - 1]] = val
        name = data.get('Channel Name')
        if name and not is_bad_row(name):
            all_rows.append(data)

    # Dedup by URL
    url_map = {}
    deduped = []
    for rd in all_rows:
        url = normalize_url(rd.get('YouTube URL'))
        if url is None:
            if rd.get('Channel Name'):
                deduped.append(rd)
            continue
        if url in url_map:
            idx = url_map[url]
            if count_fields(rd, headers) > count_fields(deduped[idx], headers):
                deduped[idx] = rd
        else:
            url_map[url] = len(deduped)
            deduped.append(rd)

    # Dedup by name
    name_map = {}
    final = []
    for rd in deduped:
        nm = str(rd.get('Channel Name', '')).strip().lower()
        clean = re.sub(r'[^\w\s]', '', nm).strip()
        if clean in name_map:
            idx = name_map[clean]
            if count_fields(rd, headers) > count_fields(final[idx], headers):
                old = final[idx]
                for h in headers:
                    if (h not in rd or rd[h] is None) and h in old and old[h] is not None:
                        rd[h] = old[h]
                final[idx] = rd
        else:
            name_map[clean] = len(final)
            final.append(rd)

    print(f"Channels: {len(final)}")

    # Apply contact updates
    updated = 0
    for rd in final:
        name = str(rd.get('Channel Name', '')).strip()
        if name in CONTACT_UPDATES:
            for field, value in CONTACT_UPDATES[name].items():
                current = rd.get(field)
                if not current or str(current).strip() == '':
                    rd[field] = value
                    updated += 1
                elif field in ('Notes', 'Other') and value not in str(current):
                    rd[field] = str(current) + '. ' + value
                    updated += 1
    print(f"Contact updates: {updated}")

    # Sort by category
    verdict_order = {'RECOMMEND': 0, 'MAYBE': 1}

    def sort_key(rd):
        cat = rd.get('Category') or 'Uncategorized'
        cat = str(cat).strip()
        if cat in ('None', ''):
            cat = 'zzz_Uncategorized'
        v = verdict_order.get(rd.get('Verdict'), 2)
        subs = parse_subs(rd.get('Subscribers'))
        return (cat.lower(), v, -subs)

    final.sort(key=sort_key)

    # Write new workbook
    new_wb = openpyxl.Workbook()
    nws = new_wb.active
    nws.title = 'Channels by Project'

    hdr_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
    hdr_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    rec_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
    may_fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
    nv_fill = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
    cat_fill = PatternFill(start_color='D6DCE4', end_color='D6DCE4', fill_type='solid')
    cat_font = Font(name='Calibri', size=11, bold=True, color='1F4E79')
    bdr = Border(left=Side(style='thin'), right=Side(style='thin'),
                 top=Side(style='thin'), bottom=Side(style='thin'))

    for col, h in enumerate(headers, 1):
        c = nws.cell(row=1, column=col, value=h)
        c.fill = hdr_fill
        c.font = hdr_font
        c.alignment = Alignment(horizontal='center', wrap_text=True)
        c.border = bdr

    cur = 2
    cur_cat = None
    num = 1

    for rd in final:
        cat = rd.get('Category') or 'Uncategorized'
        cat_d = str(cat).strip()
        if cat_d in ('None', ''):
            cat_d = 'Uncategorized'

        if cat_d != cur_cat:
            cur_cat = cat_d
            for col in range(1, len(headers) + 1):
                c = nws.cell(row=cur, column=col)
                c.fill = cat_fill
                c.border = bdr
            nws.cell(row=cur, column=3, value=f"--- {cat_d} ---").font = cat_font
            nws.cell(row=cur, column=3).fill = cat_fill
            cur += 1

        nws.cell(row=cur, column=1, value=num).border = bdr
        for col, h in enumerate(headers, 1):
            if h == '#':
                continue
            val = rd.get(h)
            c = nws.cell(row=cur, column=col, value=val)
            c.border = bdr
            c.alignment = Alignment(wrap_text=True, vertical='top')
            v = rd.get('Verdict')
            if v == 'RECOMMEND':
                c.fill = rec_fill
            elif v == 'MAYBE':
                c.fill = may_fill
            elif not v:
                c.fill = nv_fill

        num += 1
        cur += 1

    widths = {'#': 5, 'Verdict': 12, 'Channel Name': 35, 'Subscribers': 12,
              'YouTube URL': 45, 'Category': 18, 'Country': 18, 'Niche': 25,
              'Quality': 12, 'Fake Risk': 10, 'Avg Views': 15, 'View/Sub %': 10,
              'Upload Freq': 12, 'Community (1-10)': 10, 'Engagement': 12,
              'Email': 30, 'Telegram': 30, 'Instagram': 25, 'Twitter/X': 25,
              'WhatsApp': 15, 'Other': 35, 'Contact Quality': 12, 'Role': 20,
              'Notes': 50, 'Verdict Reason': 60}
    for col, h in enumerate(headers, 1):
        nws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = widths.get(h, 15)

    nws.freeze_panes = 'A2'
    nws.auto_filter.ref = f'A1:{openpyxl.utils.get_column_letter(len(headers))}{cur - 1}'

    new_wb.save('YouTubers_Report_v2_CLEANED.xlsx')

    hc = sum(1 for rd in final if any(rd.get(f) and str(rd[f]).strip() for f in ['Email', 'Telegram', 'WhatsApp']))
    print(f"With contacts: {hc}/{len(final)}")
    print(f"RECOMMEND: {sum(1 for r in final if r.get('Verdict')=='RECOMMEND')}")
    print(f"MAYBE: {sum(1 for r in final if r.get('Verdict')=='MAYBE')}")
    print("Done!")


if __name__ == '__main__':
    main()
