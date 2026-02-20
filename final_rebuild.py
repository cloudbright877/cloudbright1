"""
Full rebuild: read original, clean, enrich contacts, sort by project, save.
"""
import openpyxl
import re
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


CONTACT_UPDATES = {
    'Tabraiz Shams': {
        'Email': 'tabraizvlogs@gmail.com',
        'Telegram': 't.me/ShamsTabraiz',
        'Instagram': 'instagram.com/shams__tabraiz',
        'Contact Quality': 'HIGH',
    },
    'P4 Provider': {
        'Telegram': 't.me/P4Provider, DM: @HMTanveer',
        'WhatsApp': '0328-8855333',
        'Other': 'p4provider.com, linktr.ee/p4provider',
        'Contact Quality': 'HIGH',
    },
    'The Hobbyist Miner': {
        'Other': 'thehobbyistminer.io, Discord community',
        'Country': 'USA',
        'Contact Quality': 'MEDIUM',
    },
    'IITian in Crypto': {
        'Twitter/X': 'x.com/IITian_InCrypto',
        'Country': 'India',
    },
    'Christian Rauchenwald': {
        'Other': 'christianrauchenwald.com',
        'Contact Quality': 'MEDIUM',
    },
    'Blockchain SAGE': {
        'Telegram': 't.me/Sageblockchain',
        'Twitter/X': 'x.com/Bl0ckchainsage',
        'Contact Quality': 'MEDIUM',
    },
    'Crypto Empire': {
        'Telegram': 't.me/cryptoempireee, DM: @csell34',
        'Other': 'cryptoempireco.io, Discord',
        'Contact Quality': 'MEDIUM',
    },
    'Crypto Bee': {
        'Email': 'cryptobeeyt@gmail.com',
        'Telegram': 't.me/officialcryptobee, DM: @Jack_Berdul',
        'Contact Quality': 'HIGH',
    },
    'Darryl Boo': {
        'Telegram': 'DM: @cryptoboycash',
        'Twitter/X': 'x.com/darryl_boo',
        'Other': 'linktr.ee/darrylboo, TikTok: @darrylboocryptoboy',
        'Contact Quality': 'MEDIUM',
    },
    'Bitbull': {
        'Other': 'bitbullpro.net',
        'Country': 'Germany',
    },
    'Full Value Dan': {
        'Other': 'danteachescrypto.com, Telegram membership bot',
        'Contact Quality': 'MEDIUM',
    },
    'Your Friend Andy': {
        'Other': 'Partner: muskminers.com',
        'Country': 'USA',
    },
    'Best Passive Income Programs Online (Main)': {
        'Other': 'linktr.ee/cryptojaime, Telegram channel: 12K subs',
    },
}

NEW_CHANNELS = [
    {
        'Channel Name': 'Best Passive Income Programs Online (2nd Channel)',
        'Subscribers': '4,800',
        'YouTube URL': 'https://youtube.com/@bestpassiveincomeprogramso6200',
        'Category': 'Uniminepool', 'Country': 'USA / English',
        'Niche': 'Passive Income / Crypto', 'Quality': 'LOW-MEDIUM',
        'Fake Risk': 'LOW', 'Avg Views': '~200-500', 'Upload Freq': 'Irregular',
        'Email': 'dailyresiduals@gmail.com', 'Telegram': 't.me/jaimesorianocrypto',
        'Contact Quality': 'HIGH', 'Role': 'Active Promoter',
        'Notes': '2nd channel of Jaime Soriano (main: 19.6K subs).',
        'Verdict': 'MAYBE',
        'Verdict Reason': 'Secondary channel of established promoter.',
    },
    {
        'Channel Name': 'Mining Chamber', 'Subscribers': '56,600',
        'YouTube URL': 'https://youtube.com/c/MiningChamber',
        'Category': 'Crypto Mining', 'Country': 'English',
        'Niche': 'Crypto Mining / Hardware', 'Quality': 'MEDIUM-HIGH',
        'Fake Risk': 'LOW', 'Upload Freq': 'Regular', 'Role': 'Educator',
        'Notes': 'Mining community. Hardware, cloud mining, accessibility.',
        'Verdict': 'MAYBE',
        'Verdict Reason': 'Legitimate mining community (56K subs).',
    },
    {
        'Channel Name': 'Passive Crypto Mining', 'Subscribers': '54,000',
        'YouTube URL': 'https://youtube.com/c/PassiveCryptoMining',
        'Category': 'Crypto Mining', 'Country': 'English',
        'Niche': 'Passive Income / Mining', 'Quality': 'MEDIUM',
        'Fake Risk': 'MEDIUM', 'Upload Freq': 'Regular', 'Role': 'Promoter',
        'Notes': 'Passive income via mining. Since Jan 2021.',
        'Verdict': 'MAYBE',
        'Verdict Reason': 'Good niche alignment (54K subs, passive income).',
    },
    {
        'Channel Name': 'HashRaptor', 'Subscribers': '47,700',
        'YouTube URL': 'https://youtube.com/c/HashRaptor',
        'Category': 'Crypto Mining', 'Country': 'USA',
        'Niche': 'Mining / Smart Home / Tech', 'Quality': 'MEDIUM-HIGH',
        'Fake Risk': 'LOW', 'Upload Freq': 'Regular', 'Role': 'Educator',
        'Notes': 'Mining, crypto, home automation. Tech-savvy audience.',
        'Verdict': 'MAYBE',
        'Verdict Reason': 'Tech-focused mining channel (47K subs).',
    },
    {
        'Channel Name': 'Mad Electron Engineering', 'Subscribers': '38,900',
        'YouTube URL': 'https://youtube.com/c/MadElectronEngineering',
        'Category': 'Crypto Mining', 'Country': 'USA',
        'Niche': 'Renewable Energy / Mining', 'Quality': 'MEDIUM-HIGH',
        'Fake Risk': 'LOW', 'Upload Freq': 'Regular', 'Role': 'Educator',
        'Notes': 'Renewable energy + crypto mining experiments.',
        'Verdict': 'MAYBE',
        'Verdict Reason': 'Unique angle - renewable energy meets mining (38K).',
    },
    {
        'Channel Name': 'Cultivate Crypto', 'Subscribers': '34,700',
        'YouTube URL': 'https://youtube.com/c/CultivateCrypto',
        'Category': 'Crypto Mining', 'Country': 'Japan',
        'Niche': 'Crypto / Mining', 'Quality': 'MEDIUM',
        'Fake Risk': 'LOW', 'Upload Freq': 'Regular', 'Role': 'Educator',
        'Notes': 'Japan-based crypto mining channel.',
        'Verdict': 'MAYBE',
        'Verdict Reason': 'Japan-based mining channel (34K subs).',
    },
    {
        'Channel Name': 'ChumpChangeXD Mining & Crypto', 'Subscribers': '30,000',
        'YouTube URL': 'https://youtube.com/c/ChumpChangeXD',
        'Category': 'Crypto Mining', 'Country': 'English',
        'Niche': 'Mining / Crypto', 'Quality': 'MEDIUM',
        'Fake Risk': 'LOW', 'Upload Freq': 'Regular', 'Role': 'Educator',
        'Notes': 'Mining and crypto community-driven.',
        'Verdict': 'MAYBE',
        'Verdict Reason': 'Community-driven mining channel (30K subs).',
    },
    {
        'Channel Name': 'Genesis Mining', 'Subscribers': '26,400',
        'YouTube URL': 'https://youtube.com/c/GenesisMining',
        'Category': 'Cloud Mining', 'Country': 'Germany',
        'Niche': 'Cloud Mining / Trading', 'Quality': 'HIGH',
        'Fake Risk': 'LOW', 'Role': 'Platform (official)',
        'Notes': 'Official Genesis Mining channel. Major cloud mining provider.',
        'Verdict': 'MAYBE',
        'Verdict Reason': 'Official channel of major cloud mining brand (26K).',
    },
    {
        'Channel Name': 'Crypto Mining Belgium', 'Subscribers': '22,600',
        'YouTube URL': 'https://youtube.com/c/CryptoMiningBelgium',
        'Category': 'Crypto Mining', 'Country': 'Belgium',
        'Niche': 'Mining / Passive Income', 'Quality': 'MEDIUM-HIGH',
        'Fake Risk': 'LOW', 'Upload Freq': 'Regular', 'Role': 'Educator',
        'Notes': '240 videos since 2016. Passive income with crypto.',
        'Verdict': 'MAYBE',
        'Verdict Reason': 'Long-running Belgian mining channel (22K, since 2016).',
    },
    {
        'Channel Name': 'The Mining KiiNG', 'Subscribers': '21,200',
        'YouTube URL': 'https://youtube.com/c/TheMiningKiiNG',
        'Category': 'Crypto Mining', 'Country': 'USA',
        'Niche': 'GPU Mining / Education', 'Quality': 'MEDIUM',
        'Fake Risk': 'LOW', 'Upload Freq': 'Regular', 'Role': 'Educator',
        'Notes': 'GPU mining fundamentals, overclocking, BIOS modding.',
        'Verdict': 'MAYBE',
        'Verdict Reason': 'Technical mining educator (21K subs).',
    },
    {
        'Channel Name': 'Understanding Crypto - Herve', 'Subscribers': '13,700',
        'YouTube URL': 'https://youtube.com/c/UnderstandingCryptoHerve',
        'Category': 'Crypto Education', 'Country': 'UAE',
        'Niche': 'Crypto / Trading / Education', 'Quality': 'MEDIUM',
        'Fake Risk': 'LOW', 'Upload Freq': 'Regular', 'Role': 'Educator',
        'Notes': 'Tutorials, crypto research and TA. UAE-based.',
        'Verdict': 'MAYBE',
        'Verdict Reason': 'UAE-based crypto educator (13K subs).',
    },
    {
        'Channel Name': 'DXB Crypto', 'Subscribers': '10,800',
        'YouTube URL': 'https://youtube.com/c/DXBCrypto',
        'Category': 'Crypto', 'Country': 'UAE',
        'Niche': 'Crypto / Community', 'Quality': 'MEDIUM',
        'Fake Risk': 'LOW', 'Upload Freq': 'Regular', 'Role': 'Community',
        'Notes': 'Dubai-based crypto community channel.',
        'Verdict': 'MAYBE',
        'Verdict Reason': 'Dubai-based crypto community (10K subs).',
    },
]


def is_bad_row(name_str):
    if not name_str:
        return True
    name = str(name_str).strip()
    if re.match(r'^\[[\d.]+\]', name):
        return True
    if re.match(r'^\[AD:\d+\]', name):
        return True
    if 'LINE 48' in name:
        return True
    if name.startswith('---'):
        return True
    if name == '' or name == 'None':
        return True
    return False


def normalize_url(url):
    if not url:
        return None
    return str(url).strip().lower().replace('https://', '').replace('http://', '').rstrip('/')


def parse_subs(s):
    if not s:
        return 0
    s = str(s).replace(',', '').replace('+', '').strip()
    s = re.sub(r'[^\d]', '', s.split('.')[0] if '.' not in s else s)
    try:
        return int(s)
    except:
        return 0


def count_fields(data, headers):
    return sum(1 for h in headers if data.get(h) and str(data[h]).strip())


def main():
    print("=== STEP 1: Load original ===")
    wb = openpyxl.load_workbook('YouTubers_Report_v2_CLEANED.xlsx')
    ws = wb.active
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    print(f"Original: {ws.max_row - 1} rows")

    # Read and filter
    all_rows = []
    for row in range(2, ws.max_row + 1):
        data = {}
        for col in range(1, len(headers) + 1):
            val = ws.cell(row=row, column=col).value
            if val is not None:
                data[headers[col - 1]] = val
        name = data.get('Channel Name')
        if name and not is_bad_row(name):
            all_rows.append(data)

    print(f"After bad row filter: {len(all_rows)}")

    # Dedup by URL
    url_map = {}
    deduped = []
    for rd in all_rows:
        url = normalize_url(rd.get('YouTube URL'))
        if url is None:
            if rd.get('Channel Name'):
                deduped.append(rd)
            continue
        if url in url_map:
            idx = url_map[url]
            if count_fields(rd, headers) > count_fields(deduped[idx], headers):
                deduped[idx] = rd
        else:
            url_map[url] = len(deduped)
            deduped.append(rd)

    print(f"After URL dedup: {len(deduped)}")

    # Dedup by name
    name_map = {}
    final = []
    for rd in deduped:
        name = str(rd.get('Channel Name', '')).strip().lower()
        clean = re.sub(r'[^\w\s]', '', name).strip()
        if clean in name_map:
            idx = name_map[clean]
            if count_fields(rd, headers) > count_fields(final[idx], headers):
                old = final[idx]
                for h in headers:
                    if (h not in rd or rd[h] is None) and h in old and old[h] is not None:
                        rd[h] = old[h]
                final[idx] = rd
        else:
            name_map[clean] = len(final)
            final.append(rd)

    print(f"After name dedup: {len(final)}")

    # Add new channels (skip dupes)
    existing_urls = {normalize_url(rd.get('YouTube URL')) for rd in final if rd.get('YouTube URL')}
    existing_names = {re.sub(r'[^\w\s]', '', str(rd.get('Channel Name', '')).strip().lower()).strip() for rd in final}
    added = 0
    for ch in NEW_CHANNELS:
        cn = re.sub(r'[^\w\s]', '', str(ch['Channel Name']).strip().lower()).strip()
        url = normalize_url(ch.get('YouTube URL'))
        if cn in existing_names or (url and url in existing_urls):
            continue
        final.append(ch)
        existing_names.add(cn)
        if url:
            existing_urls.add(url)
        added += 1
    print(f"Added {added} new channels -> total: {len(final)}")

    print("\n=== STEP 2: Enrich contacts ===")
    updated = 0
    for rd in final:
        name = str(rd.get('Channel Name', '')).strip()
        if name in CONTACT_UPDATES:
            for field, value in CONTACT_UPDATES[name].items():
                current = rd.get(field)
                if not current or str(current).strip() == '':
                    rd[field] = value
                    updated += 1
    print(f"Applied {updated} contact field updates")

    print("\n=== STEP 3: Sort by Category ===")
    verdict_order = {'RECOMMEND': 0, 'MAYBE': 1}

    def norm_cat(cat):
        if not cat or str(cat).strip() in ('None', ''):
            return 'zzz_Uncategorized'
        return str(cat).strip()

    def sort_key(rd):
        cat = norm_cat(rd.get('Category'))
        v = verdict_order.get(rd.get('Verdict'), 2)
        subs = parse_subs(rd.get('Subscribers'))
        return (cat.lower(), v, -subs)

    final.sort(key=sort_key)

    print("\n=== STEP 4: Write new workbook ===")
    new_wb = openpyxl.Workbook()
    nws = new_wb.active
    nws.title = 'Channels by Project'

    # Styles
    hdr_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
    hdr_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    rec_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
    may_fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
    nv_fill = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
    cat_fill = PatternFill(start_color='D6DCE4', end_color='D6DCE4', fill_type='solid')
    cat_font = Font(name='Calibri', size=11, bold=True, color='1F4E79')
    bdr = Border(left=Side(style='thin'), right=Side(style='thin'),
                 top=Side(style='thin'), bottom=Side(style='thin'))

    for col, h in enumerate(headers, 1):
        c = nws.cell(row=1, column=col, value=h)
        c.fill = hdr_fill
        c.font = hdr_font
        c.alignment = Alignment(horizontal='center', wrap_text=True)
        c.border = bdr

    cur = 2
    cur_cat = None
    num = 1

    for rd in final:
        cat = rd.get('Category') or 'Uncategorized'
        cat_d = str(cat).strip()
        if cat_d in ('None', ''):
            cat_d = 'Uncategorized'

        if cat_d != cur_cat:
            cur_cat = cat_d
            for col in range(1, len(headers) + 1):
                c = nws.cell(row=cur, column=col)
                c.fill = cat_fill
                c.border = bdr
            nws.cell(row=cur, column=3, value=f"--- {cat_d} ---").font = cat_font
            nws.cell(row=cur, column=3).fill = cat_fill
            cur += 1

        nws.cell(row=cur, column=1, value=num).border = bdr
        for col, h in enumerate(headers, 1):
            if h == '#':
                continue
            val = rd.get(h)
            c = nws.cell(row=cur, column=col, value=val)
            c.border = bdr
            c.alignment = Alignment(wrap_text=True, vertical='top')
            v = rd.get('Verdict')
            if v == 'RECOMMEND':
                c.fill = rec_fill
            elif v == 'MAYBE':
                c.fill = may_fill
            elif not v:
                c.fill = nv_fill

        num += 1
        cur += 1

    widths = {'#': 5, 'Verdict': 12, 'Channel Name': 35, 'Subscribers': 12,
              'YouTube URL': 45, 'Category': 18, 'Country': 18, 'Niche': 25,
              'Quality': 12, 'Fake Risk': 10, 'Avg Views': 15, 'View/Sub %': 10,
              'Upload Freq': 12, 'Community (1-10)': 10, 'Engagement': 12,
              'Email': 30, 'Telegram': 30, 'Instagram': 25, 'Twitter/X': 25,
              'WhatsApp': 15, 'Other': 35, 'Contact Quality': 12, 'Role': 20,
              'Notes': 50, 'Verdict Reason': 60}
    for col, h in enumerate(headers, 1):
        nws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = widths.get(h, 15)

    nws.freeze_panes = 'A2'
    nws.auto_filter.ref = f'A1:{openpyxl.utils.get_column_letter(len(headers))}{cur - 1}'

    out = 'YouTubers_Report_v2_CLEANED.xlsx'
    new_wb.save(out)
    print(f"Saved: {out}")

    # Stats
    cats = {}
    for rd in final:
        c = str(rd.get('Category') or 'Uncategorized').strip()
        if c in ('None', ''):
            c = 'Uncategorized'
        cats[c] = cats.get(c, 0) + 1

    print(f"\n=== FINAL: {len(final)} channels, {len(cats)} categories ===")
    for c, n in sorted(cats.items(), key=lambda x: -x[1]):
        print(f"  {c}: {n}")

    hc = sum(1 for rd in final if any(rd.get(f) and str(rd[f]).strip() for f in ['Email', 'Telegram', 'WhatsApp']))
    print(f"\nWith contacts: {hc}/{len(final)}")
    print(f"RECOMMEND: {sum(1 for r in final if r.get('Verdict')=='RECOMMEND')}")
    print(f"MAYBE: {sum(1 for r in final if r.get('Verdict')=='MAYBE')}")


if __name__ == '__main__':
    main()
